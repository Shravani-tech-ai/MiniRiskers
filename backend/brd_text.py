from io import BytesIO

import fitz


def extract_text_from_bytes(file_bytes: bytes, filename: str) -> str:
    lower_name = (filename or "").lower()

    if lower_name.endswith(".pdf"):
        document = fitz.open(stream=file_bytes, filetype="pdf")
        pages = []

        for page in document:
            text = page.get_text("text")

            if text and text.strip():
                pages.append(text.strip())

        document.close()
        return "\n\n".join(pages)

    if lower_name.endswith((".txt", ".md")):
        return file_bytes.decode("utf-8", errors="replace")

    if lower_name.endswith(".docx"):
        try:
            from docx import Document
        except ImportError:
            raise ValueError(
                "DOCX support requires python-docx. Run: pip install python-docx"
            )

        document = Document(BytesIO(file_bytes))
        paragraphs = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        table_lines = []

        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]

                if cells:
                    table_lines.append(" | ".join(cells))

        return "\n".join(paragraphs + table_lines)

    raise ValueError(
        "Unsupported file type. Upload PDF, DOCX, TXT, or MD."
    )
