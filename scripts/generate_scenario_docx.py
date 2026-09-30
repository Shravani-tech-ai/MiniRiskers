"""Generate Architecture/scenario.docx — end-to-end MiniRiskers usage guide."""

from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt
except ImportError as error:
    raise SystemExit(
        "python-docx is required. Install with: pip install python-docx"
    ) from error

ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "Architecture" / "scenario.docx"


def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def build_document() -> Document:
    doc = Document()

    title = doc.add_heading("MiniRiskers — Complete System Usage Scenario", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        "This document walks through how MiniRiskers is used in a typical bank "
        "financial-crime risk assessment workflow. It covers every role, the "
        "five-stage assessment pipeline, governance controls, and in-app notifications."
    )

    add_heading(doc, "1. Purpose", 1)
    doc.add_paragraph(
        "MiniRiskers is a workbench for assessing the AML / financial-crime risk "
        "of launching a new banking product or changing an existing one. The system "
        "combines structured intake, quantitative risk scoring, regulatory evidence "
        "(RAG over Indian banking rules), AI-drafted assessment, human analyst review, "
        "and committee sign-off — with a tamper-evident audit trail throughout."
    )

    add_heading(doc, "2. Roles and Accounts", 1)
    add_bullets(
        doc,
        [
            "Business Owner — proposes changes, completes intake, submits for review, "
            "and provides evidence for approval conditions.",
            "Risk Analyst — runs the risk pipeline, reviews/overrides AI ratings, "
            "and verifies condition evidence.",
            "Risk Committee — makes final approve / approve-with-conditions / defer / reject decisions.",
            "Auditor — read-only access to all requests, audit trails, and examiner exports.",
            "Admin — superset access for support and testing.",
        ],
    )
    doc.add_paragraph(
        "Development seed accounts (see README): business_owner, risk_analyst, "
        "risk_committee, auditor, admin — each with role-specific passwords."
    )

    add_heading(doc, "3. End-to-End Scenario", 1)

    add_heading(doc, "Step 1 — Business Owner creates a change request", 2)
    add_bullets(
        doc,
        [
            "Log in as Business Owner and open Dashboard → New Request.",
            "Enter title, description, product type, business unit, and customer segment.",
            "The system assigns a unique request number (e.g. CR-2025-001).",
            "Notification: Business Owner receives “Request CR-XXXX created”.",
            "Complete structured intake: product, customer profile, geography, "
            "transactions, channels, vendors, and existing controls.",
            "Optional: upload a BRD and run AI extraction to pre-fill intake fields.",
            "Notification: “BRD extracted” when AI extraction completes.",
        ],
    )

    add_heading(doc, "Step 2 — Submit for Risk Analyst review", 2)
    add_bullets(
        doc,
        [
            "Intake must be 100% complete before submission.",
            "Click Submit for Analyst Review on the assessment page.",
            "Status moves from DRAFT to SUBMITTED; SLA clock starts (48-hour target).",
            "Notification: all Risk Analysts (or the assigned analyst) receive "
            "“New request CR-XXXX submitted”.",
        ],
    )

    add_heading(doc, "Step 3 — Risk Analyst runs the assessment pipeline", 2)
    add_bullets(
        doc,
        [
            "Log in as Risk Analyst. Open Assessments or Dashboard queue.",
            "Run in order: Generate Risk Factors → Calculate Risk → "
            "Generate Regulatory Evidence → Generate AI Assessment.",
            "The system scores 7 risk categories, applies control adjustments, "
            "retrieves RBI/PMLA/FIU regulatory passages, and drafts an AI recommendation.",
            "SLA: Risk assessment stage target is 12 hours.",
            "Notification at 75% (9 hours): “Risk assessment due within 3h on CR-XXXX”.",
            "Notification if overdue: “Risk assessment overdue on CR-XXXX”.",
        ],
    )

    add_heading(doc, "Step 4 — Analyst review and override policy", 2)
    add_bullets(
        doc,
        [
            "Analyst reads AI output and accepts or overrides the residual rating.",
            "Overrides require justification; large downgrades escalate to committee.",
            "Notification (if escalated): Committee receives “Escalated override on CR-XXXX”.",
            "Submit analyst review → stage moves to COMMITTEE_REVIEW.",
            "Notification: Committee receives “Request CR-XXXX ready for committee”.",
        ],
    )

    add_heading(doc, "Step 5 — Committee decision", 2)
    add_bullets(
        doc,
        [
            "Log in as Risk Committee. Open requests in COMMITTEE_REVIEW stage.",
            "Record decision: APPROVE, APPROVE_WITH_CONDITIONS, DEFER, or REJECT.",
            "For APPROVE_WITH_CONDITIONS, add conditions with due dates.",
            "Notification: Business Owner receives decision notification with request number.",
            "If DEFER to Business Owner: intake reopens, SLA pauses.",
            "If DEFER to Analyst: request returns for reassessment.",
        ],
    )

    add_heading(doc, "Step 6 — Approval conditions (if applicable)", 2)
    add_bullets(
        doc,
        [
            "Business Owner submits evidence for each open condition.",
            "Notification: Analyst receives “Evidence submitted for CR-XXXX”.",
            "Analyst verifies or rejects evidence (maker-checker: cannot verify own submission).",
            "Notification: Business Owner receives verified or rejected updates.",
            "When all conditions are verified, status becomes CONDITIONS_MET.",
            "Overdue conditions trigger “Condition overdue on CR-XXXX” notifications.",
        ],
    )

    add_heading(doc, "Step 7 — Audit and export", 2)
    add_bullets(
        doc,
        [
            "Auditor or Analyst exports examiner pack (PDF/JSON/CSV) from assessment page.",
            "Audit hash chain verifies tamper-evidence per request.",
            "Analytics page shows portfolio SLA metrics and cycle-time breakdown.",
        ],
    )

    add_heading(doc, "4. Notification Reference", 1)
    doc.add_paragraph(
        "Notifications appear in the bell icon (top bar). A red dot indicates unread items. "
        "Click a notification to mark it read and navigate to the relevant request."
    )

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = table.rows[0].cells
    headers[0].text = "Event"
    headers[1].text = "Recipient"
    headers[2].text = "Example"

    rows = [
        ("Request created", "Business Owner", "Request CR-2025-001 created"),
        ("Request submitted", "Risk Analyst", "New request CR-2025-001 submitted"),
        ("Risk assessment SLA warning", "Risk Analyst", "Risk assessment due within 3h"),
        ("Risk assessment overdue", "Risk Analyst", "Risk assessment overdue on CR-XXXX"),
        ("Analyst review submitted", "Risk Committee", "Ready for committee"),
        ("Override escalated", "Risk Committee", "Escalated override on CR-XXXX"),
        ("Committee decision", "Business Owner", "Approved / Rejected / Deferred"),
        ("Condition opened", "Business Owner", "Condition opened on CR-XXXX"),
        ("Evidence submitted", "Risk Analyst", "Evidence submitted for CR-XXXX"),
        ("Condition overdue", "Business Owner", "Condition overdue on CR-XXXX"),
        ("Methodology pending", "Risk Committee", "Methodology v1.1 pending approval"),
        ("Overall SLA at risk / breached", "Analyst + Committee", "48h SLA warning"),
    ]
    for event, recipient, example in rows:
        row = table.add_row().cells
        row[0].text = event
        row[1].text = recipient
        row[2].text = example

    add_heading(doc, "5. Methodology Governance", 1)
    add_bullets(
        doc,
        [
            "Risk Analyst proposes methodology changes on /methodology.",
            "Committee approves or rejects (maker-checker: proposer cannot approve).",
            "Active methodology version drives scoring weights, override policy, and SLAs.",
            "Notification: Committee notified when a version is submitted for approval.",
        ],
    )

    add_heading(doc, "6. SLA Summary", 1)
    add_bullets(
        doc,
        [
            "Overall intake-to-decision: 48 hours (pauses when returned to Business Owner).",
            "Awaiting analyst: 8 hours.",
            "Risk assessment: 12 hours.",
            "Analyst review: 12 hours.",
            "Committee review: 16 hours.",
            "At-risk threshold: 75% of each SLA target.",
        ],
    )

    add_heading(doc, "7. Getting Started (Local)", 1)
    add_bullets(
        doc,
        [
            "Backend: uvicorn backend.main:app --reload (port 8000).",
            "Frontend: npm run dev in frontend/ (port 5173).",
            "Log in with a seed account and follow the scenario above.",
            "Use the bell icon to track workflow events and SLA reminders in real time.",
        ],
    )

    return doc


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
